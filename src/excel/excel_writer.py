import os
import shutil
import openpyxl
from datetime import datetime, time

class ExcelWriter:
    """
    Non-destructive Excel Writer that populates counts into copies of existing
    traffic count templates while preserving all formatting, formulas, and merged cells.
    """
    def __init__(self, template_path: str):
        self.template_path = template_path
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template not found at {template_path}")
        self.layout_cache = self._inspect_template()

    def _inspect_template(self):
        wb = openpyxl.load_workbook(self.template_path, data_only=False)
        layout = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            if sheet_name.upper() == 'MAP' or sheet.max_row < 7:
                continue

            # Find movement blocks in row 4
            movement_blocks = {}
            for col in range(1, sheet.max_column + 1):
                val = str(sheet.cell(row=4, column=col).value or "")
                if "DIRECTION:" in val.upper():
                    dir_code = val.split(":")[-1].strip().upper()
                    # Find starting column of this block (usually 7 columns to the left of the DIRECTION text or aligned with time)
                    # For Site 15: DIRECTION is at col 9 (block starts at col 2), col 25 (starts at 18), col 41 (starts at 34)
                    start_col = col - 7 if col >= 8 else 2
                    movement_blocks[dir_code] = start_col

            # Map time rows
            time_rows = {}
            for r in range(7, sheet.max_row + 1):
                cell_val = sheet.cell(row=r, column=1).value
                if cell_val is not None:
                    if isinstance(cell_val, (datetime, time)):
                        t_str = cell_val.strftime("%H:%M:%S")
                    else:
                        t_str = str(cell_val).strip()
                    # Standardize "8:00:00" -> "08:00:00"
                    if len(t_str.split(":")[0]) == 1:
                        t_str = "0" + t_str
                    time_rows[t_str[:5]] = r  # Key by "HH:MM"

            layout[sheet_name] = {
                "movement_blocks": movement_blocks,
                "time_rows": time_rows
            }
        return layout

    def populate_counts(self, counts_matrix: dict, output_path: str):
        """
        counts_matrix structure:
        {
           'Arm A': {
               'A-B': {
                   '08:00': {'Car/Jeep/Van': 5, 'Two Wheelers (White Plate)': 20, ...},
                   '08:15': {...}
               }
           }
        }
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        shutil.copy2(self.template_path, output_path)
        
        wb = openpyxl.load_workbook(output_path)

        col_offsets = {
            "Two Wheelers (White Plate)": 0,
            "Two Wheelers (Taxi)": 1,
            "Car/Jeep/Van": 2,
            "Autorickshaw - 3wh": 3,
            "Bus - City Bus (BMTC)": 4,
            "Bus - (KSRTC) BUS": 5,
            "Bus - Other Buses": 6,
            "Bus - Mini/Midi Bus": 7,
            "Goods Auto/LCV": 8,
            "Other Trucks": 9,
            "Agricultural Tractor/Trailer": 10,
            "Cycle": 11,
            "PBS": 12,
            "Others": 13,
        }

        for sheet_name, movements in counts_matrix.items():
            if sheet_name not in wb.sheetnames:
                continue
            sheet = wb[sheet_name]
            sheet_info = self.layout_cache.get(sheet_name, {})
            blocks = sheet_info.get("movement_blocks", {})
            time_rows = sheet_info.get("time_rows", {})

            for mov_code, time_data in movements.items():
                start_col = blocks.get(mov_code.upper())
                if not start_col:
                    # Fallback for single block or matching movement
                    for b_code, b_col in blocks.items():
                        if mov_code.upper() in b_code or b_code in mov_code.upper():
                            start_col = b_col
                            break
                if not start_col:
                    continue

                for time_bucket, class_counts in time_data.items():
                    bucket_key = time_bucket[:5]
                    row_idx = time_rows.get(bucket_key)
                    if not row_idx:
                        continue

                    for class_name, count in class_counts.items():
                        offset = col_offsets.get(class_name)
                        if offset is not None:
                            target_cell = sheet.cell(row=row_idx, column=start_col + offset)
                            target_cell.value = count

        wb.save(output_path)
        print(f"Successfully populated counts into: {output_path}")
