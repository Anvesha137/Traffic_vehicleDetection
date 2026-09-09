import os
import cv2
import openpyxl
import pandas as pd

WORKSPACE = r"C:\Users\Admin\Desktop\Traffic"
VIDEO_FILE = os.path.join(WORKSPACE, "E_City_Phase1_Dmart_veerasandra_8 to 11 (1).avi")
SITE15_EXCEL = os.path.join(WORKSPACE, "Site 15 - Veerasandra_Main_Road.xlsx")
TEMP_EXCEL = os.path.join(WORKSPACE, "Data Entry Temp.xlsx")

def inspect_video():
    print("="*60)
    print("VIDEO INSPECTION")
    print("="*60)
    if not os.path.exists(VIDEO_FILE):
        print(f"Error: Video file not found at {VIDEO_FILE}")
        return
    
    cap = cv2.VideoCapture(VIDEO_FILE)
    if not cap.isOpened():
        print("Error: Failed to open video.")
        return
    
    width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
    height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration_sec = total_frames / fps if fps > 0 else 0
    duration_min = duration_sec / 60
    
    print(f"File: {os.path.basename(VIDEO_FILE)}")
    print(f"Resolution: {width} x {height}")
    print(f"FPS: {fps:.2f}")
    print(f"Total Frames: {total_frames}")
    print(f"Duration: {duration_min:.2f} minutes ({duration_sec:.1f} seconds)")
    
    ret, frame = cap.read()
    if ret:
        out_dir = os.path.join(WORKSPACE, "outputs")
        os.makedirs(out_dir, exist_ok=True)
        frame_path = os.path.join(out_dir, "sample_frame.jpg")
        cv2.imwrite(frame_path, frame)
        print(f"Sample frame saved to: {frame_path}")
    cap.release()

def inspect_excel(file_path, label):
    print("\n" + "="*60)
    print(f"EXCEL AUDIT: {label} ({os.path.basename(file_path)})")
    print("="*60)
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path, data_only=False)
    print(f"Sheet names: {wb.sheetnames}")
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n--- Sheet: {sheet_name} (Max Row: {sheet.max_row}, Max Col: {sheet.max_column}) ---")
        
        # Read first 12 rows to examine metadata and column headers
        rows = list(sheet.iter_rows(values_only=True, max_row=12))
        for r_idx, row in enumerate(rows, start=1):
            non_empty = [(c_idx, val) for c_idx, val in enumerate(row, start=1) if val is not None]
            if non_empty:
                val_strs = [f"Col{c}:{str(v)[:25]}" for c, v in non_empty[:10]]
                print(f"Row {r_idx:2d}: {', '.join(val_strs)}")
                if len(non_empty) > 10:
                    print(f"        ... and {len(non_empty)-10} more cols: {', '.join([f'Col{c}:{str(v)[:25]}' for c, v in non_empty[10:20]])}")

        # Check time interval rows (e.g. rows 7 to 35)
        time_rows = []
        for r in range(6, min(sheet.max_row+1, 40)):
            val_col_a = sheet.cell(row=r, column=1).value
            if val_col_a is not None:
                time_rows.append((r, str(val_col_a)))
        print(f"Time bucket rows found (Col 1): {time_rows[:10]} ... total {len(time_rows)}")

if __name__ == "__main__":
    inspect_video()
    inspect_excel(SITE15_EXCEL, "Site 15 TMC Template")
    inspect_excel(TEMP_EXCEL, "Data Entry Temp Template")
